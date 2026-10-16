import ast
import base64
import datetime as dt
import hashlib
import html
import io
import json
from pathlib import Path
import unittest
import xml.etree.ElementTree as ET
import zipfile
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from bom_excel_preview import workbook_preview_payload, build_bom_workbook_viewer_html, _display

ROOT = Path(__file__).resolve().parents[1]


def fixture_bytes():
    book = Workbook()
    first = book.active
    first.title = '完整BOM'
    first.append(['物料编码', '型号', '封装', '原始隐藏列', '成本', '公式'])
    for i in range(1, 224):
        first.append([f'00{i:04d}', f'FRC0402J{i} TS', 201, '不自动取消隐藏', i / 100, None])
    first['C2'].number_format = '0000'
    first['E2'].number_format = '0.0000'
    first['F2'] = '=SUM(E2:E3)'
    first['A1'].font = Font(bold=True, color='173453')
    first['A1'].fill = PatternFill('solid', fgColor='EAF1FD')
    first.column_dimensions['B'].width = 32
    first.column_dimensions['D'].hidden = True
    first.row_dimensions[5].hidden = True
    first['XFD1048576'].fill = PatternFill('solid', fgColor='FFFFFF')
    second = book.create_sheet('其他工作表')
    second.append(['说明', '日期', '百分比'])
    second.append(['<img src=x onerror=alert(1)></script><script>alert(2)</script>', dt.datetime(2026, 10, 16, 12, 30), .125])
    second['C2'].number_format = '0.0%'
    second['A2'].hyperlink = 'https://example.invalid/private-data'
    secret = book.create_sheet('原始隐藏表')
    secret['A1'] = '隐藏内容不应自动出现在预览'
    secret.sheet_state = 'hidden'
    book.create_sheet('空工作表')
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def footer_builder():
    source = (ROOT / 'component_matcher.py').read_text(encoding='utf-8')
    node = next(n for n in ast.parse(source).body if isinstance(n, ast.FunctionDef) and n.name == 'build_bom_download_footer_html')
    namespace = {'base64': base64, 'hashlib': hashlib, 'html': html, 'json': json,
                 'clean_text': lambda v: str(v or '').strip(),
                 'MEMBER_AUTH_OUTER_SHELL_ORIGIN': 'https://fruition-component.pages.dev'}
    exec(compile(ast.Module(body=[node], type_ignores=[]), 'footer-test', 'exec'), namespace)
    return namespace['build_bom_download_footer_html']


class BomExcelPreviewTests(unittest.TestCase):
    def test_display_rounding_matches_excel_not_binary_float_or_bankers_rounding(self):
        self.assertEqual(_display(2.675, '0.00'), '2.68')
        self.assertEqual(_display(-2.675, '0.00'), '-2.68')
        self.assertEqual(_display(.02675, '0.00%'), '2.68%')
        self.assertEqual(_display(1234.565, '#,##0.00'), '1,234.57')
        self.assertEqual(_display(.125, 'General'), '0.125')

    def test_whole_workbook_not_filtered_page_and_original_bytes_unchanged(self):
        raw = fixture_bytes()
        before = hashlib.sha256(raw).hexdigest()
        payload = workbook_preview_payload(raw, 'BOM_匹配后.xlsx', '其他工作表')
        self.assertEqual([s['name'] for s in payload['sheets']], ['完整BOM', '其他工作表', '空工作表'])
        self.assertEqual(payload['initial'], '其他工作表')
        self.assertEqual(len(payload['sheets'][0]['rows']), 223)
        self.assertEqual(payload['sheets'][0]['rows'][-1]['r'], 224)
        self.assertEqual(payload['sheets'][0]['rows'][-1]['c'][0][0], '000223')
        self.assertEqual([c['label'] for c in payload['sheets'][0]['columns']], ['A', 'B', 'C', 'E', 'F'])
        self.assertEqual(payload['hiddenSheets'], 1)
        self.assertEqual(payload['hiddenRanges'], 2)
        self.assertEqual(base64.b64decode(payload['xlsx']), raw)
        self.assertEqual(hashlib.sha256(raw).hexdigest(), before)
        self.assertEqual(payload['sheets'][0]['rows'][1]['c'][2][0], '0201')
        self.assertEqual(payload['sheets'][0]['rows'][1]['c'][3][0], '0.0100')
        self.assertIn('font-weight:700', payload['styles'][payload['sheets'][0]['rows'][0]['c'][0][1]])
        self.assertEqual(payload['sheets'][1]['rows'][1]['c'][1][0], '2026-10-16 12:30:00')
        self.assertEqual(payload['sheets'][1]['rows'][1]['c'][2][0], '12.5%')

    def test_formula_is_not_executed_and_cells_are_not_html(self):
        raw = fixture_bytes()
        payload = workbook_preview_payload(raw, 'report</script><script>alert(7)</script>.xlsx')
        self.assertEqual(payload['formulas'], 1)
        self.assertEqual(payload['uncached'], 1)
        self.assertEqual(payload['sheets'][0]['rows'][1]['c'][-1][0], '=SUM(E2:E3)')
        viewer, reason = build_bom_workbook_viewer_html(raw, 'report</script><script>alert(7)</script>.xlsx')
        self.assertEqual(reason, '')
        self.assertNotIn('</script><script>', viewer)
        self.assertIn('\\u003c/script\\u003e', viewer)
        self.assertNotIn('https://example.invalid/private-data', viewer)
        self.assertIn('td.textContent=cell[0]', viewer)
        self.assertNotIn('.innerHTML', viewer)
        self.assertIn("connect-src 'none'", viewer)
        self.assertNotIn('__PAYLOAD__', viewer)

    def test_invalid_or_large_file_degrades_without_truncating_or_disabling_save(self):
        self.assertEqual(build_bom_workbook_viewer_html(b'not-xlsx', 'bad.xlsx')[0], '')
        with patch('bom_excel_preview.MAX_CELLS', 10):
            viewer, reason = build_bom_workbook_viewer_html(fixture_bytes(), 'large.xlsx')
        self.assertEqual(viewer, '')
        self.assertIn('下载', reason)
        footer = footer_builder()(b'not-xlsx', 'bad.xlsx')
        self.assertIn('browserDownload(blob)', footer)
        self.assertIn('预览暂不可用', footer)
        self.assertIn(base64.b64encode(b'not-xlsx').decode(), footer)

    def test_footer_has_two_actions_private_blob_viewer_and_initial_sheet(self):
        raw = fixture_bytes()
        footer = footer_builder()(raw, 'matched.xlsx', initial_sheet='其他工作表', bridge_channel='qa-channel')
        self.assertIn('在线打开', footer)
        self.assertIn('下载 Excel', footer)
        self.assertIn('rel="noopener noreferrer"', footer)
        self.assertIn('type: "text/html;charset=utf-8"', footer)
        self.assertIn('fruition-file-save', footer)
        self.assertNotIn('__PREVIEW_', footer)
        self.assertIn('showSaveFilePicker', footer)

    def test_saved_formula_values_are_read_without_recalculation(self):
        raw = fixture_bytes()
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(output, 'w') as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == 'xl/worksheets/sheet1.xml':
                    xml = ET.fromstring(data)
                    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                    xml.find('.//s:c[@r="F2"]/s:v', ns).text = '0.03'
                    data = ET.tostring(xml)
                target.writestr(info, data)
        payload = workbook_preview_payload(output.getvalue(), 'cached.xlsx')
        self.assertEqual(payload['formulas'], 1)
        self.assertEqual(payload['uncached'], 0)
        self.assertEqual(payload['sheets'][0]['rows'][1]['c'][-1][0], '0.03')
        self.assertEqual(payload['sheets'][0]['rows'][1]['c'][-1][2], '=SUM(E2:E3)')

    def test_preview_failure_cannot_break_original_download(self):
        builder = footer_builder()
        with patch.dict('sys.modules', {'bom_excel_preview': None}):
            footer = builder(b'original-export', 'result.xlsx')
        self.assertIn('browserDownload(blob)', footer)
        self.assertIn(base64.b64encode(b'original-export').decode(), footer)
        self.assertIn('预览暂不可用', footer)


if __name__ == '__main__':
    unittest.main()
