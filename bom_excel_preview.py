"""Private, self-contained, read-only viewer of the exact matched XLSX bytes.

No workbook writes, formula execution, shared cache, public URLs or third-party
viewer. Unsupported/oversized input fails closed without affecting the download.
"""
from __future__ import annotations

import base64
import datetime as dt
from decimal import Decimal, ROUND_HALF_UP, localcontext
import io
import json
import math
import numbers
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET

MAX_INPUT_BYTES = 16 * 1024 * 1024
MAX_UNPACKED_BYTES = 128 * 1024 * 1024
MAX_SHEET_XML_BYTES = 32 * 1024 * 1024
MAX_CELLS = 500_000
MAX_COLUMNS = 256
MAX_ROWS = 100_000
MAX_PAYLOAD_CHARS = 20 * 1024 * 1024


class PreviewUnavailable(ValueError):
    pass


def _json(value):
    return (json.dumps(value, ensure_ascii=False, separators=(',', ':'))
            .replace('&', '\\u0026').replace('<', '\\u003c').replace('>', '\\u003e')
            .replace('\u2028', '\\u2028').replace('\u2029', '\\u2029'))


def _tag(node):
    return node.tag.rsplit('}', 1)[-1]


def _sheet_paths(archive):
    relationships = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
    targets = {}
    for node in relationships:
        if node.get('TargetMode') == 'External':
            continue
        target = node.get('Target', '')
        path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
        if path.startswith('xl/') and '..' not in path.split('/'):
            targets[node.get('Id')] = path
    book = ET.fromstring(archive.read('xl/workbook.xml'))
    result = {}
    for node in book.iter():
        if _tag(node) != 'sheet':
            continue
        relationship = next((v for k, v in node.attrib.items() if k.endswith('}id')), '')
        if relationship in targets:
            result[node.get('name', '')] = targets[relationship]
    return result


def _metadata(archive, path):
    if archive.getinfo(path).file_size > MAX_SHEET_XML_BYTES:
        raise PreviewUnavailable('工作表较大，请下载 Excel 查看完整内容。')
    hidden_rows, hidden_columns, widths, formulas = set(), set(), {}, {}
    last_row = last_col = physical_cells = 0
    from openpyxl.utils.cell import coordinate_to_tuple
    with archive.open(path) as source:
        root = None
        for event, node in ET.iterparse(source, events=('start', 'end')):
            if root is None:
                root = node
            if event != 'end':
                continue
            tag = _tag(node)
            if tag == 'col':
                left, right = int(node.get('min', '1')), int(node.get('max', '1'))
                for col in range(max(1, left), min(MAX_COLUMNS, right) + 1):
                    if node.get('hidden') in ('1', 'true'):
                        hidden_columns.add(col)
                    try:
                        width = float(node.get('width', '14'))
                        if math.isfinite(width):
                            widths[col] = max(70, min(360, round(width * 7 + 12)))
                    except (ValueError, OverflowError):
                        pass
                node.clear()
            elif tag == 'c':
                physical_cells += 1
                if physical_cells > MAX_CELLS:
                    raise PreviewUnavailable('工作簿较大，在线预览暂不支持；请下载 Excel 查看完整内容。')
                formula = next((n for n in node if _tag(n) == 'f'), None)
                has_value = any(_tag(n) == 'v' and n.text is not None for n in node)
                has_inline = any(_tag(n) == 't' and n.text for n in node.iter())
                coordinate = node.get('r')
                if coordinate and (has_value or has_inline or formula is not None):
                    row, col = coordinate_to_tuple(coordinate)
                    last_row, last_col = max(last_row, row), max(last_col, col)
                    if formula is not None:
                        formulas[coordinate] = '=' + formula.text if formula.text else '公式（需在 Excel 中计算）'
                node.clear()
            elif tag == 'row':
                if node.get('hidden') in ('1', 'true'):
                    hidden_rows.add(int(node.get('r', '0')))
                root.clear()
    if last_row > MAX_ROWS or last_col > MAX_COLUMNS or last_row * last_col > MAX_CELLS:
        raise PreviewUnavailable('工作表范围较大，请下载 Excel 查看完整内容；预览不会截断结果。')
    return {'rows': last_row, 'cols': last_col, 'hidden_rows': hidden_rows,
            'hidden_columns': hidden_columns, 'widths': widths, 'formulas': formulas}


def _rounded(value, precision, grouped=False):
    number = Decimal(str(value))
    with localcontext() as context:
        context.prec = max(32, len(number.as_tuple().digits) + max(0, number.adjusted()) + precision + 4)
        rounded = number.quantize(Decimal(1).scaleb(-precision), rounding=ROUND_HALF_UP)
        return format(rounded, (',' if grouped else '') + f'.{precision}f')


def _display(value, number_format='General'):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, dt.datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S') if value.time() != dt.time() or re.search(r'[hs]', number_format, re.I) else value.strftime('%Y-%m-%d')
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, numbers.Real) and math.isfinite(value):
        code = number_format.split(';', 1)[0]
        bare = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.', '', code)
        decimals = re.search(r'\.([0#]+)', bare)
        precision = min(12, len(decimals.group(1))) if decimals else 0
        if '%' in bare:
            return _rounded(Decimal(str(value)) * 100, precision) + '%'
        if re.fullmatch('0{2,}', bare) and value == int(value):
            return str(int(value)).zfill(min(50, len(bare)))
        if decimals or '#,##' in bare:
            return _rounded(value, precision, '#,##' in bare)
        return str(int(value)) if value == int(value) else str(value)
    return str(value)


def _style(cell):
    result = []
    if cell.font and cell.font.bold:
        result.append('font-weight:700')
    for source, prop in [(cell.font.color if cell.font else None, 'color'),
                         (cell.fill.fgColor if cell.fill and cell.fill.patternType == 'solid' else None, 'background')]:
        if source is not None and source.type == 'rgb' and isinstance(source.rgb, str) and re.fullmatch('[0-9A-Fa-f]{8}', source.rgb):
            result.append(prop + ':#' + source.rgb[-6:])
    if cell.alignment and cell.alignment.horizontal in ('left', 'center', 'right'):
        result.append('text-align:' + cell.alignment.horizontal)
    return ';'.join(result)


def workbook_preview_payload(data_bytes, filename, initial_sheet=''):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import get_column_letter
    if not data_bytes or len(data_bytes) > MAX_INPUT_BYTES:
        raise PreviewUnavailable('文件较大或尚未生成，请使用下载 Excel 查看。')
    with zipfile.ZipFile(io.BytesIO(data_bytes)) as archive:
        infos = archive.infolist()
        if (len(infos) > 4000 or sum(i.file_size for i in infos) > MAX_UNPACKED_BYTES
                or any(i.file_size > MAX_SHEET_XML_BYTES for i in infos if i.filename.endswith(('.xml', '.rels')))):
            raise PreviewUnavailable('工作簿较大，请下载 Excel 查看完整内容。')
        paths = _sheet_paths(archive)
        metadata = {name: _metadata(archive, path) for name, path in paths.items()
                    if '/worksheets/' in path}
    if sum(m['rows'] * m['cols'] for m in metadata.values()) > MAX_CELLS:
        raise PreviewUnavailable('工作簿较大，请下载 Excel 查看完整内容；预览不会截断结果。')
    book = load_workbook(io.BytesIO(data_bytes), read_only=True, data_only=True, keep_links=False)
    sheets, styles, style_ids = [], [''], {'': 0}
    hidden_sheets = hidden_ranges = formula_count = uncached_formulas = 0
    try:
        for ws in book.worksheets:
            if ws.sheet_state != 'visible':
                hidden_sheets += 1
                continue
            meta = metadata.get(ws.title)
            if meta is None:
                raise PreviewUnavailable('该工作簿暂不支持在线预览，请下载 Excel 查看。')
            hidden_ranges += len(meta['hidden_rows']) + len(meta['hidden_columns'])
            columns = [i for i in range(1, meta['cols'] + 1) if i not in meta['hidden_columns']]
            sheet = {'name': ws.title, 'columns': [{'label': get_column_letter(i), 'width': meta['widths'].get(i, 150)} for i in columns], 'rows': []}
            if meta['rows'] and meta['cols']:
                ws.reset_dimensions()
                for index, cells in enumerate(ws.iter_rows(min_row=1, max_row=meta['rows'], max_col=meta['cols']), 1):
                    if index in meta['hidden_rows']:
                        continue
                    row = []
                    for col in columns:
                        cell = cells[col - 1]
                        coordinate = f'{get_column_letter(col)}{index}'
                        formula = meta['formulas'].get(coordinate, '')
                        value = cell.value
                        if formula:
                            formula_count += 1
                            if value is None:
                                uncached_formulas += 1
                                value = formula
                        text = _display(value, cell.number_format)
                        css = _style(cell) if getattr(cell, 'has_style', False) else ''
                        if css not in style_ids:
                            if len(styles) < 1024:
                                style_ids[css] = len(styles)
                                styles.append(css)
                            else:
                                css = ''
                        row.append([text, style_ids[css], formula] if text or css or formula else None)
                    sheet['rows'].append({'r': index, 'c': row})
            sheets.append(sheet)
    finally:
        book.close()
    if not sheets:
        raise PreviewUnavailable('没有可见的数据工作表，请下载 Excel 查看原工作簿。')
    payload = {'filename': str(filename or 'BOM_匹配后.xlsx'), 'initial': str(initial_sheet or ''),
               'sheets': sheets, 'styles': styles, 'hiddenSheets': hidden_sheets,
               'hiddenRanges': hidden_ranges, 'formulas': formula_count, 'uncached': uncached_formulas,
               'xlsx': base64.b64encode(data_bytes).decode('ascii')}
    if len(_json(payload)) > MAX_PAYLOAD_CHARS:
        raise PreviewUnavailable('工作簿内容较大，请下载 Excel 查看完整内容；预览不会截断结果。')
    return payload


VIEWER_HTML = r'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="referrer" content="no-referrer"><meta http-equiv="Content-Security-Policy" content="default-src 'none'; script-src 'unsafe-inline'; style-src 'unsafe-inline'; img-src data:; connect-src 'none'; base-uri 'none'; form-action 'none'"><title>Excel 在线预览</title><style>
*{box-sizing:border-box}body{margin:0;font:14px/1.5 system-ui,'Microsoft YaHei',sans-serif;color:#173453;background:#f3f6fc;height:100dvh;display:flex;flex-direction:column}header{display:flex;align-items:center;gap:20px;padding:16px 22px;background:#fff;border-bottom:1px solid #d7e0ef}h1{font-size:20px;line-height:1.4;margin:0;overflow-wrap:anywhere}header p{margin:4px 0 0;font-size:12px;color:#59708f}header>div{flex:1;min-width:0}button,input,select{font:inherit;border:1px solid #bdcde4;border-radius:6px;min-height:40px;background:white;color:#173453;padding:7px 12px}button{cursor:pointer}button:disabled{opacity:.45;cursor:default}button:focus-visible,input:focus-visible,select:focus-visible{outline:3px solid #2158e6;outline-offset:2px}.download{background:#1651eb;color:white;border-color:#1651eb;white-space:nowrap}.tools{display:flex;align-items:center;gap:12px;padding:10px 22px;flex-wrap:wrap}.tools label{display:flex;gap:7px;align-items:center}input{width:250px}#summary{margin-left:auto;color:#536d8f;font-size:12px}#notice{font-size:12px;color:#5c6f88;margin:0;padding:0 22px 9px}#table-wrap{min-height:0;flex:1;overflow:auto;background:white;border-top:1px solid #d7e0ef;border-bottom:1px solid #d7e0ef}table{table-layout:fixed;border-collapse:separate;border-spacing:0;width:max-content;min-width:100%}th,td{border-right:1px solid #dce3ef;border-bottom:1px solid #dce3ef;vertical-align:top;padding:8px 10px;white-space:pre-wrap;overflow-wrap:anywhere;font-weight:400;line-height:1.6}thead th{position:sticky;top:0;z-index:3;background:#edf2fa;color:#4c6485;text-align:center;font-size:12px}th.row-index{position:sticky;left:0;z-index:2;width:55px;min-width:55px;background:#edf2fa;text-align:center;color:#4c6485;font-size:12px}thead .row-index{z-index:4}td.formula-only{color:#906424;background:#fff9e9}footer{background:white;flex-shrink:0}#sheets{display:flex;overflow-x:auto;gap:6px;padding:9px 18px;border-bottom:1px solid #d7e0ef}#sheets button{white-space:nowrap;min-width:80px}#sheets button[aria-pressed=true]{background:#e7efff;border-color:#1651eb;color:#1651eb;font-weight:700}.pages{display:flex;gap:10px;align-items:center;justify-content:center;padding:9px 14px}.empty{padding:36px;text-align:center;color:#637898}@media(max-width:600px){header{padding:10px 12px;gap:10px}h1{font-size:16px}header p{font-size:10px}.tools{padding:8px 12px;gap:8px}.tools label{font-size:12px}input{width:min(220px,65vw)}#summary{margin-left:0;width:100%}#notice{padding:0 12px 7px;font-size:10px}th,td{font-size:12px;padding:7px}.pages{gap:7px;font-size:12px}#sheets{padding:7px 10px}button,input,select{min-height:44px}.pages input{width:58px!important}}
</style></head><body><header><div><h1 id="filename"></h1><p>Excel 在线预览 · 只读 · 来自完整匹配文件，不受原页面筛选和分页影响</p></div><button id="download" class="download">下载 Excel</button></header><section class="tools"><label>查找<input id="search" type="search" placeholder="在当前工作表中查找" autocomplete="off"></label><label>每页<select id="page-size"><option value="100">100行</option><option value="200">200行</option><option value="500">500行</option></select></label><span id="summary" role="status" aria-live="polite"></span></section><p id="notice"></p><main id="table-wrap" tabindex="0" aria-label="工作表数据"></main><footer><nav id="sheets" aria-label="工作表"></nav><div class="pages"><button id="prev">上一页</button><label>第 <input id="page" type="number" min="1" value="1" style="width:70px"> 页</label><span id="page-count"></span><button id="next">下一页</button></div></footer><script>
(()=>{'use strict';const book=__PAYLOAD__,el=id=>document.getElementById(id);let current=Math.max(0,book.sheets.findIndex(s=>s.name===book.initial)),page=1,size=100,rows=[];
el('filename').textContent=book.filename;document.title=book.filename+' · Excel 在线预览';
const notices=['本页只读显示表格数据和常用格式；图表、宏及复杂排版请下载 Excel 查看。'];if(book.formulas)notices.push(book.uncached?'有 '+book.uncached+' 个公式没有保存计算值，已显示原式；本页不执行公式。':'公式显示文件中已保存的计算值，本页不重新计算。');if(book.hiddenSheets||book.hiddenRanges)notices.push('保留原文件的隐藏工作表/行列设置，不自动取消隐藏。');el('notice').textContent=notices.join(' ');
function filter(){const q=el('search').value.toLocaleLowerCase();rows=book.sheets[current].rows.filter(r=>!q||r.c.some(c=>c&&c[0].toLocaleLowerCase().includes(q)));page=1;render();}
function render(){const sheet=book.sheets[current],pages=Math.max(1,Math.ceil(rows.length/size));page=Math.max(1,Math.min(pages,Math.trunc(page)||1));el('summary').textContent=sheet.name+' · '+rows.length+' / '+sheet.rows.length+' 行 · '+sheet.columns.length+' 列 · '+book.sheets.length+' 个可见工作表';el('page').value=page;el('page').max=pages;el('page-count').textContent='/ '+pages;el('prev').disabled=page<=1;el('next').disabled=page>=pages;
 for(const b of el('sheets').children)b.setAttribute('aria-pressed',String(Number(b.dataset.sheet)===current));
 const wrap=el('table-wrap');wrap.replaceChildren();if(!sheet.columns.length){const empty=document.createElement('p');empty.className='empty';empty.textContent='空工作表';wrap.append(empty);return;}
 const table=document.createElement('table'),cols=document.createElement('colgroup'),indexCol=document.createElement('col');indexCol.style.width='55px';cols.append(indexCol);for(const c of sheet.columns){const col=document.createElement('col');col.style.width=c.width+'px';cols.append(col);}table.append(cols);const head=document.createElement('thead'),tr=document.createElement('tr'),corner=document.createElement('th');corner.className='row-index';corner.textContent='行';tr.append(corner);for(const c of sheet.columns){const th=document.createElement('th');th.textContent=c.label;tr.append(th);}head.append(tr);table.append(head);const body=document.createElement('tbody');
 for(const row of rows.slice((page-1)*size,page*size)){const tr=document.createElement('tr'),th=document.createElement('th');th.className='row-index';th.scope='row';th.textContent=row.r;tr.append(th);for(const cell of row.c){const td=document.createElement('td');if(cell){td.textContent=cell[0];td.style.cssText=book.styles[cell[1]]||'';if(cell[2]){td.title=cell[2];if(cell[0]===cell[2])td.classList.add('formula-only');}}tr.append(td);}body.append(tr);}table.append(body);wrap.append(table);wrap.scrollTop=0;
 for(const b of el('sheets').children)b.setAttribute('aria-pressed',String(Number(b.dataset.sheet)===current));}
book.sheets.forEach((s,i)=>{const b=document.createElement('button');b.textContent=s.name;b.dataset.sheet=i;b.onclick=()=>{current=i;el('search').value='';filter();};el('sheets').append(b);});
el('search').oninput=filter;el('page-size').onchange=()=>{size=Number(el('page-size').value);page=1;render();};el('prev').onclick=()=>{page--;render();};el('next').onclick=()=>{page++;render();};el('page').onchange=()=>{page=Number(el('page').value)||1;render();};
el('download').onclick=()=>{const binary=atob(book.xlsx),bytes=Uint8Array.from(binary,c=>c.charCodeAt(0)),url=URL.createObjectURL(new Blob([bytes],{type:'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'})),a=document.createElement('a');a.href=url;a.download=book.filename;document.body.append(a);a.click();a.remove();setTimeout(()=>URL.revokeObjectURL(url),30000);};filter();})();
</script></body></html>'''


def build_bom_workbook_viewer_html(data_bytes, filename, initial_sheet=''):
    """Return (self-contained viewer, reason if unavailable); never break saving."""
    try:
        payload = workbook_preview_payload(data_bytes, filename, initial_sheet)
        return VIEWER_HTML.replace('__PAYLOAD__', _json(payload)), ''
    except PreviewUnavailable as error:
        return '', str(error)
    except Exception:
        return '', '该文件暂不能在线预览，请下载 Excel 查看；原文件不会被修改。'
